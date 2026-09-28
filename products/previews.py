"""販売ページ用のプレビュー画像を、実際の商品ファイルから作る。

使い方: python3 products/previews.py
必要なもの: LibreOffice（Calc）、PyMuPDF、BIZ UD フォント（見た目を本番に近づけるため）
出力: site/assets/img/*.png
"""
import shutil
import subprocess
import tempfile
from pathlib import Path

import pymupdf
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
DIST = ROOT / 'products' / 'dist'
IMG = ROOT / 'site' / 'assets' / 'img'
XLSX = DIST / '個別支援計画 文例つき様式セット.xlsx'
BOOK = DIST / '個別支援計画 文例全集.pdf'


def to_pdf(xlsx: Path, workdir: Path) -> Path:
    subprocess.run(['soffice', '--headless', '--norestore', '--convert-to', 'pdf', '--outdir', str(workdir), str(xlsx)],
                   check=True, capture_output=True, timeout=240)
    return workdir / (xlsx.stem + '.pdf')


def save_page(pdf: Path, marker, out: Path, *, dpi=110):
    """marker（ページ番号、またはページ内の文字列）で選んだページを、描画範囲に合わせて切り取って保存する。"""
    doc = pymupdf.open(pdf)
    if isinstance(marker, int):
        page = doc[marker]
    else:
        page = next(p for p in doc if marker in p.get_text())
    rect = pymupdf.Rect()
    for b in page.get_text('blocks'):
        rect |= pymupdf.Rect(b[:4])
    for d in page.get_drawings():
        rect |= d['rect']
    rect = (rect + (-12, -12, 12, 12)) & page.rect if not rect.is_empty else page.rect
    page.get_pixmap(dpi=dpi, clip=rect).save(out)
    print(out)


def zoom_plan(pdf: Path, out: Path):
    """「課題を選ぶ」列から支援内容の列までの、文例が入った3行ぶんを拡大して保存する。"""
    page = pymupdf.open(pdf)[0]
    top = page.search_for('③ 課題を選ぶ')[0]
    left = page.search_for('▼ ここで操作します')[0]
    right = page.search_for('達成')[0]
    bottom = page.search_for('感覚過敏（音・触覚など）がある')[0]
    clip = pymupdf.Rect(left.x0 - 4, top.y0 - 14, right.x0 - 3, bottom.y1 + 34)
    page.get_pixmap(dpi=240, clip=clip).save(out)
    print(out)


def sheet_pdf(wb_path: Path, keep, workdir: Path, prepare=None) -> Path:
    """keep に挙げたシートだけを残して PDF にする（数式が参照するシートも keep に入れること）。"""
    wb = load_workbook(wb_path)
    for name in list(wb.sheetnames):
        if name not in keep:
            del wb[name]
    if prepare:
        prepare(wb)
    src = workdir / f'p{abs(hash(tuple(keep)))}.xlsx'
    wb.save(src)
    return to_pdf(src, workdir)


def main():
    IMG.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        data = ['文例一覧', '対象データ']

        def fill(wb):
            ws = wb['計画書（文例入り）']
            ws['A4'] = '小学校低学年'
            ws['A7'] = '電車、ブロック'
            ws['A10'] = '活動の切り替えが苦手'
            ws['A11'] = '友だちとの関わりが少ない'
            ws['A12'] = '感覚過敏（音・触覚など）がある'
            ws.print_area = 'A1:I17'
            ws.page_setup.fitToHeight = 1
            if 'モニタリング記録' in wb.sheetnames:
                wb['モニタリング記録']['D6'] = '達成'
                wb['モニタリング記録']['D7'] = '一部達成'
                wb['モニタリング記録'].page_setup.fitToHeight = 1

        def list_area(wb):
            ws = wb['文例一覧']
            ws.print_area = 'A1:F6'
            ws.page_setup.fitToHeight = 1

        def one_page(name):
            def f(wb):
                wb[name].page_setup.fitToHeight = 1
            return f

        save_page(sheet_pdf(XLSX, ['計画書（文例入り）'] + data, tmp, fill), '▼ ここで操作します', IMG / 'excel-plan.png', dpi=130)
        zoom_plan(sheet_pdf(XLSX, ['計画書（文例入り）'] + data, tmp, fill), IMG / 'excel-plan-zoom.png')
        save_page(sheet_pdf(XLSX, ['計画書（文例入り）', 'モニタリング記録'] + data, tmp, fill), 'モニタリング記録', IMG / 'excel-monitoring.png', dpi=110)
        save_page(sheet_pdf(XLSX, ['記入例'], tmp, one_page('記入例')), 0, IMG / 'excel-example.png', dpi=120)
        save_page(sheet_pdf(XLSX, ['文例一覧'], tmp, list_area), 0, IMG / 'excel-list.png', dpi=110)
        free = ROOT / 'site' / 'files' / 'kobetsu-shien-keikaku-yoshiki.xlsx'
        save_page(sheet_pdf(free, ['個別支援計画書'], tmp, one_page('個別支援計画書')), 0, IMG / 'free-plan.png', dpi=100)
    save_page(BOOK, 0, IMG / 'book-cover.png', dpi=70)
    save_page(BOOK, 1, IMG / 'book-toc.png', dpi=70)
    save_page(BOOK, 16, IMG / 'book-issue.png', dpi=90)


if __name__ == '__main__':
    main()
