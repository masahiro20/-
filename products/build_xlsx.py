"""有料の「文例つき様式セット」と、無料配布の白紙様式（Excel）を作る。

使い方:
    node scripts/export-data.mjs      # 文例データを products/build/data.json に書き出す
    python3 products/build_xlsx.py    # products/dist/ と site/files/ に xlsx を出力
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
DATA = json.loads((ROOT / 'products' / 'build' / 'data.json').read_text(encoding='utf-8'))
DOMAINS = {d['id']: d for d in DATA['domains']}
AGES = DATA['ages']
ISSUES = DATA['issues']
CONFIG = DATA['config']

DIST = ROOT / 'products' / 'dist'
FREE = ROOT / 'site' / 'files'

FONT = 'BIZ UDPGothic'  # Windows では「BIZ UDPゴシック」として表示される
AI = '1F3B60'
HEAD_FILL = PatternFill('solid', fgColor='EEF1F5')
INPUT_FILL = PatternFill('solid', fgColor='FFF6CC')
NOTE_FILL = PatternFill('solid', fgColor='FBEFEC')
WHITE = PatternFill('solid', fgColor='FFFFFF')
THIN = Side(style='thin', color='2B2F33')
HAIR = Side(style='thin', color='C9C4BA')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
SOFT_BOX = Border(left=HAIR, right=HAIR, top=HAIR, bottom=HAIR)
DIAG = Border(left=THIN, right=THIN, top=THIN, bottom=THIN, diagonal=THIN, diagonalUp=True)

WRAP_TOP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(wrap_text=True, vertical='center', horizontal='center')

N_ROWS = 5            # 本人支援の行数
DATA_LAST = 300       # 文例一覧で参照する最終行（事業所オリジナルの文例を追加できる余白込み）
SHEET_DATA = '文例一覧'
SHEET_AGE = '対象データ'


def font(size=10, bold=False, color='1C2126'):
    return Font(name=FONT, size=size, bold=bold, color=color)


def put(ws, ref, value, *, size=10, bold=False, color='1C2126', fill=None, align=WRAP_TOP, border=BOX):
    c = ws[ref]
    c.value = value
    c.font = font(size, bold, color)
    c.alignment = align
    if fill:
        c.fill = fill
    if border:
        c.border = border
    return c


def box_range(ws, rng, border=BOX):
    for row in ws[rng]:
        for c in row:
            c.border = border


def merge(ws, rng, value=None, **kw):
    ws.merge_cells(rng)
    box_range(ws, rng, kw.get('border', BOX))
    first = rng.split(':')[0]
    return put(ws, first, value, **kw)


def supports_text(issue):
    doms = []
    for s in issue['supports']:
        for d in s['domains']:
            if d not in doms:
                doms.append(d)
    body = '\n'.join('・' + s['text'] for s in issue['supports'])
    return body + '\n【5領域】' + '／'.join(DOMAINS[d]['name'] for d in doms), '／'.join(DOMAINS[d]['name'] for d in doms)


# ── データシート ──
DATA_COLS = [
    ('領域', 12), ('課題', 26), ('アセスメントで確かめたいこと', 34), ('支援目標（文例1）', 38), ('支援目標（文例2）', 38),
    ('支援内容（5領域つき）', 52), ('関連する5領域', 18), ('留意事項', 30), ('家族支援の目標', 34), ('家族支援の内容', 36),
    ('モニタリング（達成）', 36), ('モニタリング（継続）', 36), ('長期目標の書き出し', 30),
]
COL = {name: chr(ord('A') + i) for i, (name, _) in enumerate(DATA_COLS)}


def rng(col_name):
    c = COL[col_name]
    return f"'{SHEET_DATA}'!${c}$2:${c}${DATA_LAST}"


def lookup(col_name, key_ref):
    return f"INDEX({rng(col_name)},MATCH({key_ref},{rng('課題')},0))"


def build_data_sheet(wb):
    ws = wb.create_sheet(SHEET_DATA)
    for i, (name, width) in enumerate(DATA_COLS):
        col = chr(ord('A') + i)
        put(ws, f'{col}1', name, bold=True, color='FFFFFF', fill=PatternFill('solid', fgColor=AI), align=CENTER, border=SOFT_BOX)
        ws.column_dimensions[col].width = width
    for r, issue in enumerate(ISSUES, start=2):
        sup, doms = supports_text(issue)
        row = [
            DOMAINS[issue['domain']]['name'], issue['label'], '\n'.join('・' + a for a in issue['assess']),
            issue['shortGoals'][0], issue['shortGoals'][1] if len(issue['shortGoals']) > 1 else issue['shortGoals'][0],
            sup, doms, issue['note'], issue['familyGoal'], issue['family'],
            issue['monitoring']['done'], issue['monitoring']['cont'], issue['longGoal'],
        ]
        for i, v in enumerate(row):
            put(ws, f'{chr(ord("A") + i)}{r}', v, border=SOFT_BOX)
        ws.row_dimensions[r].height = 96
    first_free = len(ISSUES) + 2
    put(ws, f'A{first_free}', '（ここから下に、事業所オリジナルの文例を追加できます）', color='8A6D00', fill=INPUT_FILL, border=SOFT_BOX)
    ws.merge_cells(f'A{first_free}:{chr(ord("A") + len(DATA_COLS) - 1)}{first_free}')
    for r in range(first_free + 1, first_free + 21):
        for i in range(len(DATA_COLS)):
            c = ws[f'{chr(ord("A") + i)}{r}']
            c.fill = INPUT_FILL
            c.border = SOFT_BOX
            c.font = font()
            c.alignment = WRAP_TOP
    ws.freeze_panes = 'C2'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = '1:1'
    ws.auto_filter.ref = f'A1:{chr(ord("A") + len(DATA_COLS) - 1)}{len(ISSUES) + 1}'
    ws.sheet_properties.tabColor = '7A8086'
    # プルダウン用の名前（追加した文例も自動で候補に出る）
    wb.defined_names['課題リスト'] = DefinedName(
        '課題リスト', attr_text=f"OFFSET('{SHEET_DATA}'!$B$2,0,0,MAX(1,COUNTA('{SHEET_DATA}'!$B$2:$B${DATA_LAST})),1)")
    return ws


def build_age_sheet(wb):
    ws = wb.create_sheet(SHEET_AGE)
    cols = [('対象', 16), ('生活の場面', 22), ('連携先', 10), ('移行支援の目標', 40), ('移行支援の内容', 60),
            ('地域支援の目標', 40), ('地域支援の内容', 50)]
    for i, (name, width) in enumerate(cols):
        col = chr(ord('A') + i)
        put(ws, f'{col}1', name, bold=True, color='FFFFFF', fill=PatternFill('solid', fgColor=AI), align=CENTER, border=SOFT_BOX)
        ws.column_dimensions[col].width = width
    for r, a in enumerate(AGES, start=2):
        for i, v in enumerate([a['name'], a['context'], a['partner'], a['transitionGoal'], a['transition'], a['communityGoal'], a['community']]):
            put(ws, f'{chr(ord("A") + i)}{r}', v, border=SOFT_BOX)
        ws.row_dimensions[r].height = 80
    ws.sheet_properties.tabColor = '7A8086'
    return ws


def age_lookup(col, age_ref):
    return f"INDEX('{SHEET_AGE}'!${col}$2:${col}$5,MATCH({age_ref},'{SHEET_AGE}'!$A$2:$A$5,0))"


# ── 計画書シート ──
# with_helper=True のとき A〜B 列が「課題を選ぶ」「文例」の操作列、C〜I 列が様式（印刷範囲）
def build_plan_sheet(wb, title, *, with_helper, example=None):
    ws = wb.create_sheet(title)
    off = 2 if with_helper else 0
    cols = [chr(ord('A') + off + i) for i in range(7)]  # 項目, 目標, 内容, 時期, 担当, 留意, 優先
    C, Dg, E, F, G, H, I = cols
    widths = [9, 30, 48, 8, 15, 24, 6]
    for col, w in zip(cols, widths):
        ws.column_dimensions[col].width = w
    if with_helper:
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 7

    # 見出し
    ws.row_dimensions[1].height = 30
    merge(ws, f'{C}1:{Dg}1', '利用児氏名：', size=10, align=Alignment(vertical='bottom'), border=None)
    merge(ws, f'{E}1:{F}1', '個別支援計画書', size=16, bold=True, align=Alignment(horizontal='center', vertical='center'), border=None)
    merge(ws, f'{G}1:{I}1', '作成年月日：　　年　　月　　日', size=10, align=Alignment(horizontal='right', vertical='bottom'), border=None)

    heights = {3: 48, 4: 62, 5: 66, 6: 66}
    for r, h in heights.items():
        ws.row_dimensions[r].height = h
    put(ws, f'{C}3', '利用児及び家族の生活に対する意向', size=9, bold=True, fill=HEAD_FILL, align=CENTER)
    put(ws, f'{C}4', '総合的な支援の方針', size=9, bold=True, fill=HEAD_FILL, align=CENTER)
    put(ws, f'{C}5', '長期目標\n（内容・期間等）', size=9, bold=True, fill=HEAD_FILL, align=CENTER)
    put(ws, f'{C}6', '短期目標\n（内容・期間等）', size=9, bold=True, fill=HEAD_FILL, align=CENTER)
    merge(ws, f'{F}5:{G}6', '支援の標準的な\n提供時間等\n（曜日・頻度、時間）', size=9, bold=True, fill=HEAD_FILL, align=CENTER)

    row0 = 10
    rows_self = list(range(row0, row0 + N_ROWS))
    r_fam, r_tr, r_com = row0 + N_ROWS, row0 + N_ROWS + 1, row0 + N_ROWS + 2
    age_ref = '$A$4'
    likes_ref = '$A$7'

    if with_helper:
        wish = policy = long_goal = short_goal = time = None
        has_any = f'COUNTIF($A${row0}:$A${row0 + N_ROWS - 1},"?*")>0'
        first_issue = f'$A${row0}'
        policy = (f'=IF(NOT({has_any}),"",IF({likes_ref}="","本人の好きなことや得意なことを活かしながら、",'
                  f'"本人の好きなこと・得意なこと（"&{likes_ref}&"）を活かしながら、")'
                  f'&IF({first_issue}="","","特に「"&{lookup("領域", first_issue)}&"」の領域を中心に、")'
                  f'&"5領域の視点をふまえた総合的な支援を行う。家庭・"&IFERROR({age_lookup("C", age_ref)},"学校")'
                  f'&"・関係機関と連携しながら、"&IFERROR({age_lookup("B", age_ref)},"学校や家庭での生活")&"を安心して過ごせるよう支える。")')
        long_goal = (f'=IF($A${row0}="","",{lookup("長期目標の書き出し", f"$A${row0}")}&"、"'
                     f'&IF($A${row0 + 1}="","",{lookup("長期目標の書き出し", f"$A${row0 + 1}")}&"、")'
                     f'&IFERROR({age_lookup("B", age_ref)},"学校や家庭での生活")&"を自信を持って過ごすことができる。"&CHAR(10)&"（期間：1年）")')
        short_goal = (f'=IF({Dg}{row0}="","","・"&{Dg}{row0}&IF({Dg}{row0 + 1}="","",CHAR(10)&"・"&{Dg}{row0 + 1})'
                      f'&CHAR(10)&"（期間：6か月）")')
    else:
        wish = policy = long_goal = short_goal = time = None
    if example:
        wish, policy, long_goal, short_goal, time = (example[k] for k in ('wish', 'policy', 'long', 'short', 'time'))

    merge(ws, f'{Dg}3:{I}3', wish)
    merge(ws, f'{Dg}4:{I}4', policy)
    merge(ws, f'{Dg}5:{E}5', long_goal)
    merge(ws, f'{Dg}6:{E}6', short_goal)
    merge(ws, f'{H}5:{I}6', time)
    for r in (3, 4, 5, 6):
        ws[f'{C}{r}'].border = BOX
    # 太枠
    for r in range(3, 7):
        ws[f'{C}{r}'].border = Border(left=Side(style='medium'), right=THIN, top=Side(style='medium') if r == 3 else THIN,
                                      bottom=Side(style='medium') if r == 6 else THIN)

    put(ws, f'{C}8', '○支援目標及び具体的な支援内容等', bold=True, border=None, align=Alignment(vertical='bottom'))
    ws.row_dimensions[9].height = 44
    heads = ['項目', '支援目標\n（具体的な到達目標）', '支援内容\n（内容・支援の提供上のポイント・5領域との関連性等）',
             '達成\n時期', '担当者\n提供機関', '留意事項', '優先\n順位']
    for col, h in zip(cols, heads):
        put(ws, f'{col}9', h, size=9, bold=True, fill=HEAD_FILL, align=CENTER)

    for n, r in enumerate(rows_self):
        ws.row_dimensions[r].height = 124 if example else 100
        put(ws, f'{C}{r}', '本人支援', size=9, bold=True, fill=HEAD_FILL, align=CENTER)
        key = f'$A{r}'
        if with_helper:
            put(ws, f'{Dg}{r}', f'=IF({key}="","",IF($B{r}=2,{lookup("支援目標（文例2）", key)},{lookup("支援目標（文例1）", key)}))')
            put(ws, f'{E}{r}', f'=IF({key}="","",{lookup("支援内容（5領域つき）", key)})')
            put(ws, f'{F}{r}', f'=IF({key}="","",IF(COUNTIF($A${row0}:{key},"?*")=1,"3か月後","6か月後"))', align=CENTER)
            put(ws, f'{G}{r}', f'=IF({key}="","","児童指導員・保育士")')
            put(ws, f'{H}{r}', f'=IF({key}="","",{lookup("留意事項", key)})')
            put(ws, f'{I}{r}', f'=IF({key}="","",COUNTIF($A${row0}:{key},"?*"))', align=CENTER)
        else:
            for col in cols[1:]:
                put(ws, f'{col}{r}', None, align=CENTER if col in (F, I) else WRAP_TOP)
        if example and n < len(example['rows']):
            ex = example['rows'][n]
            for col, v in zip(cols[1:], ex):
                put(ws, f'{col}{r}', v, align=CENTER if col in (F, I) else WRAP_TOP)

    labels = {r_fam: '家族支援', r_tr: '移行支援', r_com: '地域支援・\n地域連携'}
    for r, lab in labels.items():
        ws.row_dimensions[r].height = 112 if r == r_fam else 84
        put(ws, f'{C}{r}', lab, size=9, bold=True, fill=HEAD_FILL, align=CENTER)
        for col in cols[1:6]:
            put(ws, f'{col}{r}', None, align=CENTER if col == F else WRAP_TOP)
        ws[f'{I}{r}'].border = DIAG
    if with_helper and not example:
        def first_nonblank(col_name):
            expr = '""'
            for r in reversed(rows_self):
                expr = f'IF($A{r}<>"",{lookup(col_name, f"$A{r}")},{expr})'
            return '=' + expr
        put(ws, f'{Dg}{r_fam}', first_nonblank('家族支援の目標'))
        put(ws, f'{E}{r_fam}', '=' + '&'.join(f'IF($A{r}="","","・"&{lookup("家族支援の内容", f"$A{r}")}&CHAR(10))' for r in rows_self))
        put(ws, f'{Dg}{r_tr}', f'=IFERROR({age_lookup("D", age_ref)},"")')
        put(ws, f'{E}{r_tr}', f'=IFERROR({age_lookup("E", age_ref)},"")')
        put(ws, f'{Dg}{r_com}', f'=IFERROR({age_lookup("F", age_ref)},"")')
        put(ws, f'{E}{r_com}', f'=IFERROR({age_lookup("G", age_ref)},"")')
        for r in (r_fam, r_tr, r_com):
            put(ws, f'{F}{r}', f'=IF({has_any},"6か月後","")', align=CENTER)
        put(ws, f'{G}{r_fam}', f'=IF({has_any},"児童発達支援管理責任者","")')
        put(ws, f'{G}{r_tr}', f'=IF({has_any},"児童発達支援管理責任者（連携先："&IFERROR({age_lookup("C", age_ref)},"在籍校")&"）","")')
        put(ws, f'{G}{r_com}', f'=IF({has_any},"児童発達支援管理責任者（連携先："&IFERROR({age_lookup("C", age_ref)},"在籍校")&"・相談支援事業所）","")')
        put(ws, f'{H}{r_fam}', f'=IF({has_any},"家族支援加算を算定する場合は、実施方法・頻度を記載する。","")')
        put(ws, f'{H}{r_com}', f'=IF({has_any},"関係機関連携加算を算定する場合は、その旨を記載する。","")')
    if example:
        for r, ex in zip((r_fam, r_tr, r_com), example['others']):
            for col, v in zip(cols[1:6], ex):
                put(ws, f'{col}{r}', v, align=CENTER if col == F else WRAP_TOP)

    foot = r_com + 2
    merge(ws, f'{C}{foot}:{E}{foot}', '提供する支援内容について、本計画書に基づき説明しました。', border=None)
    merge(ws, f'{F}{foot}:{I}{foot}', '本計画書に基づき支援の説明を受け、内容に同意しました。', border=None)
    merge(ws, f'{C}{foot + 1}:{E}{foot + 1}', '児童発達支援管理責任者氏名：', border=None)
    merge(ws, f'{F}{foot + 1}:{I}{foot + 1}', '　　年　　月　　日（保護者署名）', border=None)
    ws.row_dimensions[foot].height = 20
    ws.row_dimensions[foot + 1].height = 26

    # 操作列
    if with_helper and not example:
        put(ws, 'A1', '▼ ここで操作します', bold=True, color=AI, border=None)
        put(ws, 'A3', '① 対象', bold=True, border=None, align=Alignment(vertical='bottom'))
        put(ws, 'A4', '小学校低学年', fill=INPUT_FILL, align=Alignment(vertical='center'))
        put(ws, 'A6', '② 好きなこと・得意なこと（任意）', bold=True, border=None, align=Alignment(vertical='bottom', wrap_text=True))
        put(ws, 'A7', None, fill=INPUT_FILL)
        ws.row_dimensions[7].height = 24
        put(ws, 'A9', '③ 課題を選ぶ（上から順に）', size=9, bold=True, fill=INPUT_FILL, align=CENTER)
        put(ws, 'B9', '文例\n1/2', size=9, bold=True, fill=INPUT_FILL, align=CENTER)
        for r in rows_self:
            put(ws, f'A{r}', None, fill=INPUT_FILL, align=Alignment(vertical='center', wrap_text=True))
            put(ws, f'B{r}', 1, fill=INPUT_FILL, align=CENTER)
        put(ws, f'A{r_fam}', '← 家族支援・移行支援・地域支援は、選んだ課題と対象から自動で入ります。', size=9, color='767D84',
            border=None, align=Alignment(vertical='top', wrap_text=True))
        dv_age = DataValidation(type='list', formula1=f"'{SHEET_AGE}'!$A$2:$A$5", allow_blank=True)
        dv_issue = DataValidation(type='list', formula1='課題リスト', allow_blank=True)
        dv_issue.error = '文例一覧にある課題から選んでください。'
        dv_issue.prompt = '課題を選ぶと、支援目標・支援内容・留意事項に文例が入ります。'
        dv_issue.showInputMessage = True
        dv_ex = DataValidation(type='list', formula1='"1,2"', allow_blank=True)
        for dv in (dv_age, dv_issue, dv_ex):
            ws.add_data_validation(dv)
        dv_age.add('A4')
        dv_issue.add(f'A{row0}:A{row0 + N_ROWS - 1}')
        dv_ex.add(f'B{row0}:B{row0 + N_ROWS - 1}')

    # 印刷設定
    last = foot + 1
    ws.print_area = f'{C}1:{I}{last}'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.freeze_panes = None
    return ws


# ── モニタリング記録 ──
def build_monitoring_sheet(wb, title, *, linked):
    ws = wb.create_sheet(title)
    off = 1 if linked else 0
    cols = [chr(ord('A') + off + i) for i in range(6)]  # 項目, 目標, 評価, 理由, 方針, 備考
    C, Dg, E, F, G, H = cols
    for col, w in zip(cols, [9, 34, 11, 46, 11, 22]):
        ws.column_dimensions[col].width = w
    if linked:
        ws.column_dimensions['A'].width = 4
    ws.row_dimensions[1].height = 30
    merge(ws, f'{C}1:{Dg}1', '利用児氏名：', border=None, align=Alignment(vertical='bottom'))
    merge(ws, f'{E}1:{F}1', 'モニタリング記録', size=16, bold=True, border=None, align=Alignment(horizontal='center', vertical='center'))
    merge(ws, f'{G}1:{H}1', '実施日：　　年　　月　　日', border=None, align=Alignment(horizontal='right', vertical='bottom'))
    put(ws, f'{C}3', '計画作成日', size=9, bold=True, fill=HEAD_FILL, align=CENTER)
    merge(ws, f'{Dg}3:{Dg}3', '　　年　　月　　日')
    put(ws, f'{E}3', '実施者', size=9, bold=True, fill=HEAD_FILL, align=CENTER)
    merge(ws, f'{F}3:{H}3', None)
    ws.row_dimensions[3].height = 24

    ws.row_dimensions[5].height = 40
    for col, h in zip(cols, ['項目', '支援目標', '評価', '評価の理由・本人の様子', '今後の\n方針', '備考']):
        put(ws, f'{col}5', h, size=9, bold=True, fill=HEAD_FILL, align=CENTER)
    labels = ['本人支援'] * N_ROWS + ['家族支援', '移行支援', '地域支援・\n地域連携']
    plan = "'計画書（文例入り）'"
    dv_eval = DataValidation(type='list', formula1='"達成,一部達成,未達成"', allow_blank=True)
    dv_plan = DataValidation(type='list', formula1='"継続,変更,終了"', allow_blank=True)
    ws.add_data_validation(dv_eval)
    ws.add_data_validation(dv_plan)
    for n, lab in enumerate(labels):
        r = 6 + n
        ws.row_dimensions[r].height = 78
        put(ws, f'{C}{r}', lab, size=9, bold=True, fill=HEAD_FILL, align=CENTER)
        plan_row = 10 + n
        if linked:
            put(ws, f'{Dg}{r}', f'=IF({plan}!D{plan_row}="","",{plan}!D{plan_row})')
            if n < N_ROWS:
                key = f'{plan}!$A{plan_row}'
                put(ws, f'{F}{r}', f'=IF(OR({key}="",{E}{r}=""),"",IF({E}{r}="達成",{lookup("モニタリング（達成）", key)},{lookup("モニタリング（継続）", key)}))')
            else:
                put(ws, f'{F}{r}', None)
        else:
            put(ws, f'{Dg}{r}', None)
            put(ws, f'{F}{r}', None)
        put(ws, f'{E}{r}', None, align=CENTER)
        put(ws, f'{G}{r}', None, align=CENTER)
        put(ws, f'{H}{r}', None)
        dv_eval.add(f'{E}{r}')
        dv_plan.add(f'{G}{r}')
    last = 6 + len(labels)
    ws.row_dimensions[last].height = 70
    put(ws, f'{C}{last}', '本人・保護者\nの意見', size=9, bold=True, fill=HEAD_FILL, align=CENTER)
    merge(ws, f'{Dg}{last}:{H}{last}', None)
    ws.row_dimensions[last + 1].height = 70
    put(ws, f'{C}{last + 1}', '総合的な\n評価', size=9, bold=True, fill=HEAD_FILL, align=CENTER)
    merge(ws, f'{Dg}{last + 1}:{H}{last + 1}', None)
    merge(ws, f'{C}{last + 3}:{F}{last + 3}', '児童発達支援管理責任者氏名：', border=None)
    if linked:
        put(ws, 'A5', None, border=None)
        merge(ws, f'{C}{last + 5}:{H}{last + 5}',
              '使い方：「評価」で「達成」を選ぶと達成の文例、「一部達成」「未達成」を選ぶと継続の文例が「評価の理由」に入ります。'
              'お子さまの様子に合わせて書き直してください。', size=9, color='767D84', border=None)
    ws.print_area = f'{C}1:{H}{last + 3}'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    return ws


# ── はじめに ──
def build_intro(wb, *, paid):
    ws = wb.active
    ws.title = 'はじめに'
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 3
    ws.column_dimensions['B'].width = 6
    ws.column_dimensions['C'].width = 96
    r = 2
    put(ws, f'B{r}', CONFIG['productName'] if paid else '個別支援計画書・モニタリング記録 様式（無料版）', size=16, bold=True, color=AI, border=None)
    ws.merge_cells(f'B{r}:C{r}')
    ws.row_dimensions[r].height = 30
    r += 1
    put(ws, f'B{r}', f'放課後等デイサービス・児童発達支援 ／ 5領域対応（令和6年度報酬改定）　ver.{CONFIG["productVersion"]}', color='767D84', border=None)
    ws.merge_cells(f'B{r}:C{r}')
    r += 2

    def heading(text):
        nonlocal r
        put(ws, f'B{r}', text, size=12, bold=True, border=None)
        ws.merge_cells(f'B{r}:C{r}')
        ws.row_dimensions[r].height = 24
        r += 1

    def step(no, text, height=34):
        nonlocal r
        put(ws, f'B{r}', no, size=12, bold=True, color=AI, border=None, align=Alignment(horizontal='center', vertical='top'))
        put(ws, f'C{r}', text, border=None)
        ws.row_dimensions[r].height = height
        r += 1

    if paid:
        heading('いちばん簡単な使い方（3分）')
        step('1', '下のタブ「計画書（文例入り）」を開きます。')
        step('2', '左の黄色いセル「① 対象」で、お子さまの区分（未就学／小学校低学年／小学校高学年／中学生・高校生）を選びます。')
        step('3', '「③ 課題を選ぶ」の黄色いセルで、気になる課題をプルダウンから選びます（上から優先順位の順に、最大5つ）。\n'
                  '→ 支援目標・支援内容（5領域つき）・留意事項・家族支援・移行支援・地域支援に、文例が自動で入ります。', 48)
        step('4', '隣の「文例」で 2 を選ぶと、支援目標が別の文例に切り替わります。')
        step('5', '入った文例を、お子さまの様子に合わせて書き直します。セルに直接上書きしてかまいません\n'
                  '（上書きしたセルは自動入力が止まり、書いた文章がそのまま残ります）。', 48)
        step('6', '印刷は、A4横に収まるよう設定済みです。黄色い操作列（A・B列）は印刷されません。')
        r += 1
        heading('そのほかのタブ')
        step('・', '計画書（白紙）…… 文例を使わずに一から書くときの様式です。')
        step('・', 'モニタリング記録 …… 計画書の支援目標が自動で入ります。「評価」を選ぶと、評価の理由の文例が入ります。')
        step('・', '記入例 …… 架空のお子さまで書いた計画書の記入例です。書き方の目安にしてください。')
        step('・', f'文例一覧 …… {len(ISSUES)}課題すべての文例です。フィルターで領域ごとに絞り込めます。\n'
                  '一番下の黄色い行に、事業所オリジナルの文例を追加すると、「課題を選ぶ」のプルダウンにも出てきます。', 48)
        step('・', '対象データ …… 移行支援・地域支援の文例です。事業所に合わせて書き換えられます。')
    else:
        heading('この様式について')
        step('・', 'こども家庭庁の参考様式（令和6年度報酬改定）と同じ項目の、個別支援計画書とモニタリング記録の様式です。')
        step('・', 'A4横に収まるよう印刷設定済みです。事業所名の追加など、自由に編集してお使いください。')
        r += 1
        heading('文例つきの有料版もあります')
        step('・', f'「{CONFIG["productName"]}」では、課題をプルダウンで選ぶだけで、支援目標・支援内容（5領域つき）・留意事項・'
                  f'家族支援・移行支援の文例が自動で入ります。モニタリング記録の文例、記入例、{len(ISSUES)}課題の文例一覧つきです。', 48)
        step('・', f'くわしくは {CONFIG["siteUrl"]}/template.html をご覧ください。')

    r += 1
    heading('ご利用にあたって')
    step('・', '文例は、計画を書き始めるための下書きです。お子さまのアセスメントと、本人・家族の意向に合わせて必ず書き直してください。')
    step('・', '様式や記載方法の細かな求めは自治体（指定権者）によって異なります。最新の資料をご確認ください。')
    if paid:
        step('・', 'ご購入いただいた事業所（同じ法人の複数事業所を含みます）の中で、ご自由にお使いいただけます。'
                  'ファイルの再配布・転売・インターネット上での公開はご遠慮ください。', 48)
        step('・', 'アップデート版は、購入された方に無料でお届けします。')
    step('・', '出典：こども家庭庁 支援局障害児支援課 事務連絡（令和6年5月17日）「個別支援計画書の記載のポイント（参考様式版）」')
    r += 1
    contact = CONFIG.get('supportEmail') or f'{CONFIG["siteUrl"]}/contact.html'
    put(ws, f'B{r}', f'お問い合わせ：{contact}　／　{CONFIG["siteName"]}', color='767D84', border=None)
    ws.merge_cells(f'B{r}:C{r}')
    ws.sheet_properties.tabColor = 'B63B27'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


EXAMPLE = {
    'wish': '（本人）友だちとカードゲームがしたい。\n（保護者）好きな遊びをやめるときに怒ってしまうことが減り、学校の準備を自分でできるようになってほしい。',
    'policy': '本人の好きなこと・得意なこと（電車、カードゲーム）を活かしながら、特に「認知・行動」「人間関係・社会性」の領域を中心に、'
              '5領域の視点をふまえた総合的な支援を行う。家庭・在籍校・関係機関と連携しながら、学校や家庭での生活を安心して過ごせるよう支える。',
    'long': '見通しを持って活動を切り替えられるようになり、友だちと関わる楽しさを知り、人とのつながりを広げ、学校や家庭での生活を自信を持って過ごすことができる。\n（期間：1年）',
    'short': '・タイマーが鳴ったら、声かけ2回以内でカードゲームを終えて片付けを始めることができる。\n・職員の仲立ちのもとで、友だち1〜2人とカードゲームを10分間楽しむことができる。\n（期間：6か月）',
    'time': '月・水・金\n14:30〜17:30\n（学校休業日 10:00〜16:00）',
    'rows': [
        ['タイマーが鳴ったら、声かけ2回以内でカードゲームを終えて片付けを始めることができる。',
         '・来所時に、その日の流れと遊びの時間をスケジュール表で一緒に確認する。\n・終わりの5分前・1分前に予告し、残り時間が目で見て分かるタイマーを使う。\n・カードは「続きは次回」と決めた箱に保管し、再開できる見通しを示す。\n・切り替えられたときは、すぐに具体的に認める。\n【5領域】認知・行動／人間関係・社会性',
         '3か月後', '児童指導員（〇〇）', '好きな遊びを急に止めない。予告の言葉を職員間でそろえる。', 1],
        ['職員の仲立ちのもとで、友だち1〜2人とカードゲームを10分間楽しむことができる。',
         '・本人の好きなカードゲームを入口に、2〜3人の小集団の遊びを設定する。\n・「かして」「次はだれの番？」など、声のかけ方を職員が見本で示す。\n・順番表を使い、自分の番が目で見て分かるようにする。\n【5領域】人間関係・社会性／言語・コミュニケーション／認知・行動',
         '6か月後', '保育士（〇〇）', '負けたときの気持ちの切り替え方を、遊びの前に一緒に確認する。', 2],
        ['来所時の身支度（靴・上着・カバン）を、手順表を見ながら自分で行うことができる。',
         '・身支度の手順を写真付きの手順表で示し、1つ終わるごとに自分でチェックできるようにする。\n・できた工程を具体的に認め、職員が手伝う量を段階的に減らす。\n【5領域】健康・生活／認知・行動',
         '6か月後', '児童指導員（〇〇）', '来所直後は落ち着いてから取り組む。', 3],
    ],
    'others': [
        ['保護者が、事業所で効果のあった予告の方法を知り、家庭での切り替えの場面で使うことができる。',
         '・タイマーや予告の言葉など、効果のあった方法を連絡帳と送迎時に共有する。\n・3か月ごとに面談を行い、家庭での様子や困りごとを聞き取る。',
         '6か月後', '児童発達支援管理責任者（〇〇）', '家族支援加算（個別・月1回）の算定を想定。'],
        ['事業所で身につけた切り替えの工夫を、学校でも使うことができる。',
         '・学校で困りやすい場面を保護者・在籍校と共有し、事業所で効果のあった予告の方法を担任に伝える。\n・地域の児童館で遊ぶ活動を月1回取り入れる。',
         '6か月後', '児童発達支援管理責任者（〇〇）\n連携先：〇〇小学校', ''],
        ['在籍校と支援の方法をそろえ、学校でも安心して過ごすことができる。',
         '・在籍校の担任と学期に1回情報交換を行う。\n・相談支援専門員と支援方針を共有する。',
         '6か月後', '児童発達支援管理責任者（〇〇）\n連携先：〇〇小学校・〇〇相談支援事業所', '関係機関連携加算の算定を想定。'],
    ],
}


def build_paid():
    wb = Workbook()
    build_intro(wb, paid=True)
    plan = build_plan_sheet(wb, '計画書（文例入り）', with_helper=True)
    build_plan_sheet(wb, '計画書（白紙）', with_helper=False)
    build_monitoring_sheet(wb, 'モニタリング記録', linked=True)
    ex = build_plan_sheet(wb, '記入例', with_helper=False, example=EXAMPLE)
    ex['A1'].value = '利用児氏名：Aさん（小学3年）※架空の記入例です'
    build_data_sheet(wb)
    build_age_sheet(wb)
    plan.sheet_properties.tabColor = AI
    ex.sheet_properties.tabColor = '3F7A52'
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True  # 開いたときに必ず再計算させる
    DIST.mkdir(parents=True, exist_ok=True)
    out = DIST / f'{CONFIG["productName"]}.xlsx'
    wb.save(out)
    return out


def build_free():
    wb = Workbook()
    build_intro(wb, paid=False)
    build_plan_sheet(wb, '個別支援計画書', with_helper=False)
    build_monitoring_sheet(wb, 'モニタリング記録', linked=False)
    FREE.mkdir(parents=True, exist_ok=True)
    out = FREE / 'kobetsu-shien-keikaku-yoshiki.xlsx'
    wb.save(out)
    return out


if __name__ == '__main__':
    print(build_paid())
    print(build_free())
